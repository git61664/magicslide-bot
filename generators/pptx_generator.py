import os
import asyncio
from datetime import datetime
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

async def generate_presentation(data: dict, user_id: int) -> str:
    if not PPTX_AVAILABLE:
        # Fallback - oddiy text fayl yaratamiz
        output_path = f"temp/taqdimot_{user_id}_{int(datetime.now().timestamp())}.txt"
        os.makedirs("temp", exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(f"TAQDIMOT: {data['topic']}\n")
            f.write(f"Muallif: {data['author']}\n")
            f.write(f"Slaydlar: {data['slides']}\n")
            f.write(f"Til: {data['language']}\n\n")
            for i in range(1, data['slides'] + 1):
                f.write(f"Slayd {i}: {data['topic']} - qism {i}\n\n")
        return output_path
    
    # PPTX generatsiya
    prs = Presentation()
    
    # Title slide
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    
    title.text = data['topic']
    subtitle.text = data['author']
    
    # Content slides
    for i in range(1, data['slides']):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        title.text = f"{data['topic']} - Qism {i}"
        content.text = f"Bu yerda {i}-qism mazmuni bo'ladi."
    
    output_path = f"temp/taqdimot_{user_id}_{int(datetime.now().timestamp())}.pptx"
    os.makedirs("temp", exist_ok=True)
    prs.save(output_path)
    
    return output_path
