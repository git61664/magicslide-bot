import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import random

async def generate_referat(data: dict, user_id: int) -> str:
    doc = Document()
    
    title = doc.add_paragraph()
    title_run = title.add_run(data['topic'])
    title_run.bold = True
    title_run.font.size = Pt(16)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info = doc.add_paragraph()
    info.add_run(f"Muassasa: {data['institution']}\n")
    info.add_run(f"Muallif: {data['author']}\n")
    info.add_run(f"Til: {data['language']}")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading("Kirish", level=1)
    doc.add_paragraph(
        "Ushbu referatda mavzuning dolzarbligi, maqsadi va vazifalari yoritiladi. "
        "Tadqiqot metodologiyasi ilmiy-nazariy tahlil asosida olingan."
    )
    
    doc.add_heading("Asosiy qism", level=1)
    for i in range(1, 3):
        doc.add_heading(f"Bo'lim {i}", level=2)
        doc.add_paragraph(
            "Bu bo'limda mavzuning asosiy jihatlari, ilmiy xulosalar va tahlillar keltirilgan. "
            "Har bir mavzu turli manbalarda o'rganilgan va isbotlangan."
        )
    
    doc.add_heading("Xulosa", level=1)
    doc.add_paragraph(
        "Referatda keltirilgan ma'lumotlar asosida quyidagi xulosalarga kelish mumkin. "
        "Mavzuni chuqur o'rganish ilm-fanni rivojlantiradiga sahodida muhim ahamiyat kasb etadi."
    )
    
    doc.add_heading("Foydalanilgan adabiyotlar", level=1)
    doc.add_paragraph("1. Qanoat manbasi (2024)")
    doc.add_paragraph("2. Ilmiy tadqiqot (2024)")
    
    output_path = f"temp/referat_{user_id}_{int(datetime.now().timestamp())}.docx"
    os.makedirs("temp", exist_ok=True)
    doc.save(output_path)
    
    return output_path

async def generate_test(data: dict, user_id: int) -> str:
    doc = Document()
    
    title = doc.add_paragraph()
    title_run = title.add_run(f"TEST: {data['topic']}")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(f"Tuzuvchi: {data['author']}")
    doc.add_paragraph(f"Daraja: {data['level']}")
    doc.add_paragraph(f"Til: {data['language']}\n")
    
    for i in range(1, data['count'] + 1):
        doc.add_paragraph(f"{i}. Savol {i}")
        for j in range(1, 5):
            doc.add_paragraph(f"   {chr(96+j)}) Variant {j}", style='List Bullet')
    
    doc.add_page_break()
    doc.add_heading("JAVOBLAR KALITI", level=1)
    
    for i in range(1, data['count'] + 1):
        variant = random.choice(['a', 'b', 'c', 'd'])
        doc.add_paragraph(f"{i}. {variant.upper()}")
    
    output_path = f"temp/test_{user_id}_{int(datetime.now().timestamp())}.docx"
    os.makedirs("temp", exist_ok=True)
    doc.save(output_path)
    
    return output_path

async def generate_kurs_ishi(data: dict, user_id: int) -> str:
    doc = Document()
    
    title = doc.add_paragraph(data['topic'])
    title.style = 'Heading 1'
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    if data.get('group'):
        doc.add_paragraph(f"Kurs/guruh: {data['group']}")
    
    doc.add_paragraph(f"Til: {data['language']}\n")
    
    doc.add_heading("Mundarija", level=1)
    doc.add_paragraph("Kirish", style='List Number')
    doc.add_paragraph("Birinchi bob", style='List Number')
    doc.add_paragraph("Ikkinchi bob", style='List Number')
    doc.add_paragraph("Xulosa", style='List Number')
    doc.add_paragraph("Adabiyotlar", style='List Number')
    
    doc.add_page_break()
    
    doc.add_heading("Kirish", level=1)
    doc.add_paragraph("Kurs ishining kirish bo'limida mavzuning dolzarbligi va ahamiyati yoritiladi.")
    
    doc.add_heading("1-Bob", level=1)
    doc.add_paragraph("Asosiy konsepsiyalar va nazariy asoslari.")
    
    doc.add_heading("2-Bob", level=1)
    doc.add_paragraph("Amaliy qo'llanish va tahlillar.")
    
    doc.add_heading("Xulosa", level=1)
    doc.add_paragraph("Kurs ishida olingan xulosalar va tavsiyalar.")
    
    doc.add_heading("Adabiyotlar ro'yxati", level=1)
    doc.add_paragraph("1. Manba 1 (2024)")
    doc.add_paragraph("2. Manba 2 (2024)")
    
    output_path = f"temp/kurs_ishi_{user_id}_{int(datetime.now().timestamp())}.docx"
    os.makedirs("temp", exist_ok=True)
    doc.save(output_path)
    
    return output_path

async def generate_other_work(data: dict, user_id: int) -> str:
    wtype = data['work_type']
    doc = Document()
    
    title = doc.add_paragraph(data['topic'])
    title.style = 'Heading 1'
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(f"Til: {data['language']}\n")
    
    if wtype == "insho":
        doc.add_paragraph("Bu insho mavzuning dolzarbligi va ahamiyatini chuqur yorituvchi yaxlit matn ko'rinishida yozilgan.")
    elif wtype == "maqola":
        doc.add_heading("Annotatsiya", level=2)
        doc.add_paragraph("Qisqa annotatsiya.")
        doc.add_heading("Kalit so'zlar", level=2)
        doc.add_paragraph("so'z1, so'z2, so'z3")
    elif wtype == "tezis":
        doc.add_paragraph("Konferensiya to'plami uchun 2-3 betlik qisqa ilmiy ish.")
    elif wtype == "keys":
        doc.add_heading("Amaliy vaziyat", level=2)
        doc.add_paragraph("Vaziyat tavsifi...")
        doc.add_heading("Muammo", level=2)
        doc.add_paragraph("Muammo analizi...")
        doc.add_heading("Yechim", level=2)
        doc.add_paragraph("Taklif etilgan yechimlar...")
    elif wtype == "glossariy":
        doc.add_table(rows=2, cols=2)
        table = doc.tables[0]
        table.rows[0].cells[0].text = "Atama"
        table.rows[0].cells[1].text = "Izoh"
        for i in range(1, 11):
            row = table.add_row()
            row.cells[0].text = f"Atama {i}"
            row.cells[1].text = f"Atamaning qisqa izohlamalyasi"
    elif wtype == "krossvord":
        doc.add_paragraph("To'ldiriladigan katakcha grid")
        doc.add_heading("Savollar", level=2)
        for i in range(1, 11):
            doc.add_paragraph(f"{i}. Savol {i}")
    
    output_path = f"temp/work_{wtype}_{user_id}_{int(datetime.now().timestamp())}.docx"
    os.makedirs("temp", exist_ok=True)
    doc.save(output_path)
    
    return output_path
